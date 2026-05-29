from __future__ import annotations

import re
from dataclasses import dataclass
from io import BytesIO
from typing import BinaryIO

from openpyxl import load_workbook

from ..validation import validate_project_fields
from .project import build_llm_model_lookup, resolve_llm_model_id

HEADER_ALIASES = {
    "маркер": "marker",
    "глубина": "depth",
    "ширина": "width",
    "n-грамма": "n_gramms",
    "n_gramm": "n_gramms",
    "n_грамма": "n_gramms",
    "кол-во разделов": "count_chapters",
    "количество разделов": "count_chapters",
    "блок о авторе": "about_author",
    "о авторе": "about_author",
    "модель": "llm_model",
    "llm": "llm_model",
    "llm-модель": "llm_model",
    "llm модель": "llm_model",
    "model": "llm_model",
}


@dataclass(frozen=True)
class ParsedProjectRow:
    marker: str
    depth: list[tuple[str, int]]
    width: list[str]
    n_gramms: list[str]
    count_chapters: int
    about_author: str
    llm_model_id: int | None = None


@dataclass(frozen=True)
class RowError:
    row: int
    message: str


def effective_llm_model_id(
    row: ParsedProjectRow,
    fallback_llm_model_id: int | None,
) -> int | None:
    """Модель из строки Excel, иначе выбор из формы импорта."""
    if row.llm_model_id is not None:
        return row.llm_model_id
    return fallback_llm_model_id


def _normalize_header(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip().lower()


def parse_depth(text: str) -> list[tuple[str, int]]:
    parts = [part.strip() for part in text.split(",") if part.strip()]
    if not parts:
        raise ValueError("Глубина не может быть пустой")
    result: list[tuple[str, int]] = []
    for part in parts:
        match = re.match(r"^(.+?)\s+(\d+)$", part.strip())
        if not match:
            raise ValueError(f"Неверный формат глубины: {part!r}")
        keyword = match.group(1).strip()
        count = int(match.group(2))
        result.append((keyword, count))
    return result


def parse_list(text: str) -> list[str]:
    normalized = text.replace("\xa0", " ")
    return [item.strip() for item in normalized.split(",") if item.strip()]


def parse_workbook(
    file: BinaryIO,
    *,
    default_llm_model_id: int | None = None,
) -> tuple[list[tuple[int, ParsedProjectRow]], list[RowError]]:
    workbook = load_workbook(BytesIO(file.read()), read_only=True, data_only=True)
    sheet = workbook.active
    rows_iter = sheet.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        return [], [RowError(row=1, message="Файл пустой")]

    column_map: dict[str, int] = {}
    for index, cell in enumerate(header_row):
        key = HEADER_ALIASES.get(_normalize_header(cell))
        if key:
            column_map[key] = index

    required = {"marker", "depth", "width", "n_gramms", "count_chapters", "about_author"}
    missing = required - set(column_map)
    if missing:
        labels = ", ".join(sorted(missing))
        return [], [RowError(row=1, message=f"Отсутствуют колонки: {labels}")]

    has_model_column = "llm_model" in column_map
    llm_lookup = build_llm_model_lookup() if has_model_column else None

    parsed_rows: list[tuple[int, ParsedProjectRow]] = []
    errors: list[RowError] = []

    for row_number, row in enumerate(rows_iter, start=2):
        if not row or all(cell is None or str(cell).strip() == "" for cell in row):
            continue

        def cell_value(field: str) -> str:
            index = column_map[field]
            if index >= len(row) or row[index] is None:
                return ""
            return str(row[index]).strip()

        marker = cell_value("marker")
        if not marker:
            errors.append(RowError(row=row_number, message="Маркер не может быть пустым"))
            continue

        try:
            depth = parse_depth(cell_value("depth"))
            width = parse_list(cell_value("width"))
            n_gramms = parse_list(cell_value("n_gramms"))
            count_raw = cell_value("count_chapters")
            if not count_raw:
                raise ValueError("Количество разделов не указано")
            count_chapters = int(float(count_raw))
            about_author = cell_value("about_author")
            if not about_author:
                raise ValueError("Блок о авторе не может быть пустым")

            llm_model_id = default_llm_model_id
            if has_model_column:
                model_raw = cell_value("llm_model")
                if model_raw:
                    llm_model_id = resolve_llm_model_id(model_raw, llm_lookup)

            row_data = ParsedProjectRow(
                marker=marker,
                depth=depth,
                width=width,
                n_gramms=n_gramms,
                count_chapters=count_chapters,
                about_author=about_author,
                llm_model_id=llm_model_id,
            )
            validate_project_fields(
                marker=row_data.marker,
                depth=row_data.depth,
                width=row_data.width,
                n_gramms=row_data.n_gramms,
                count_chapters=row_data.count_chapters,
                about_author=row_data.about_author,
            )
            parsed_rows.append((row_number, row_data))
        except LookupError as exc:
            errors.append(RowError(row=row_number, message=str(exc)))
        except (ValueError, TypeError) as exc:
            errors.append(RowError(row=row_number, message=str(exc)))
        except Exception as exc:
            errors.append(RowError(row=row_number, message=str(exc)))

    workbook.close()
    return parsed_rows, errors
